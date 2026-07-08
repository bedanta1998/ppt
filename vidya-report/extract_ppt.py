import sys
try:
    from pptx import Presentation
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

def extract_text_from_ppt(ppt_path):
    prs = Presentation(ppt_path)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)
        print("\n")

if __name__ == "__main__":
    extract_text_from_ppt(sys.argv[1])
